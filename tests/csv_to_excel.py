import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "hasil_pengujian_kriptografi.csv"
OUT_DIR = ROOT / "docs"
OUT_DIR.mkdir(exist_ok=True)
OUT_PATH = OUT_DIR / "hasil_pengujian_kriptografi.xlsx"

def keterangan(entropi: float) -> str:
    if entropi >= 7.9:
        return "Entropi mendekati maksimum (8.0) - ciphertext acak sempurna"
    elif entropi >= 7.0:
        return "Entropi tinggi - ciphertext tampak acak"
    else:
        return "Entropi lebih rendah, kemungkinan karena ukuran data kecil"

if not CSV_PATH.exists():
    CSV_PATH = ROOT / "tests" / "hasil_pengujian_kriptografi.csv"

with open(CSV_PATH, newline="", encoding="utf-8") as f:
    reader = csv.DictReader(f)
    rows = list(reader)

wb = Workbook()
ws = wb.active
ws.title = "Hasil Pengujian"

headers = [
    "Nama Berkas / Skenario", "Ukuran (Bytes)", "Waktu Enkripsi (ms)",
    "Waktu Dekripsi (ms)", "Entropi Plaintext", "Entropi Ciphertext", 
    "Avalanche Effect (%)", "Keterangan"
]
ws.append(headers)

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

for row in rows:
    nama = row.get("Nama Berkas / Skenario") or row.get("Nama Berkas", "Unknown")
    ukuran = int(row.get("Ukuran (Bytes)") or row.get("Ukuran Berkas (byte)", 0))
    enc_time = float(row.get("Waktu Enkripsi (ms)", 0.0))
    dec_time = float(row.get("Waktu Dekripsi (ms)", 0.0))
    pt_entropy = float(row.get("Entropi Plaintext") or row.get("Entropi Plainteks", 0.0))
    ct_entropy = float(row.get("Entropi Ciphertext") or row.get("Entropi Cipherteks", 0.0))
    avalanche = float(row.get("Avalanche Effect (%)", 0.0))
    
    ws.append([
        nama,
        ukuran,
        enc_time,
        dec_time,
        pt_entropy,
        ct_entropy,
        avalanche,
        keterangan(ct_entropy),
    ])

for col in ws.columns:
    max_len = max(len(str(c.value or '')) for c in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

chart = BarChart()
chart.title = "Waktu Enkripsi vs Dekripsi per Skenario"
chart.y_axis.title = "Waktu (ms)"
chart.x_axis.title = "Skenario"
data = Reference(ws, min_col=3, max_col=4, min_row=1, max_row=ws.max_row)
cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
ws.add_chart(chart, "J2")

wb.save(OUT_PATH)
print(f"Selesai: Berkas Excel berhasil disimpan di: {OUT_PATH}")